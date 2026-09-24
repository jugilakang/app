import calendar
import io
import logging
import os
import random
import uuid
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import List, Literal, Optional

import jwt
from dotenv import load_dotenv
from fastapi import APIRouter, Depends, FastAPI, HTTPException, Query
from fastapi.responses import StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from passlib.context import CryptContext
from pydantic import BaseModel, Field
from starlette.middleware.cors import CORSMiddleware

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

JWT_SECRET = os.environ["JWT_SECRET"]
JWT_ALG = "HS256"
JWT_DAYS = 30
ADMIN_USERNAME = os.environ["ADMIN_USERNAME"]
ADMIN_PASSWORD = os.environ["ADMIN_PASSWORD"]

WIB = timezone(timedelta(hours=7))
GROUPS = ["Teknisi", "Besi", "Kayu", "Finishing", "Tukang", "Kuli"]

# Batas shift (menit dari tengah malam, WIB)
MORNING_END = 11 * 60 + 30      # 11:30 -> akhir shift pagi
AFTERNOON_START = 13 * 60       # 13:00 -> mulai shift siang
EVENING_REST_END = 18 * 60      # 18:00 -> lembur mulai
LUNCH_REST_HOURS = 1.5          # 11.30-13.00, dibayar bila no_rest_siang
EVENING_REST_HOURS = 1.0        # 17.00-18.00, dibayar bila no_rest_sore

pwd_ctx = CryptContext(schemes=["bcrypt"], deprecated="auto")
bearer = HTTPBearer(auto_error=False)

api_router = APIRouter(prefix="/api")


# ---------- helpers ----------

def now_wib() -> datetime:
    return datetime.now(WIB)


def today_wib() -> str:
    return now_wib().date().isoformat()


def now_iso() -> str:
    return now_wib().isoformat()


def now_hhmm() -> str:
    return now_wib().strftime("%H:%M")


def to_min(t: str) -> int:
    h, m = t.split(":")
    return int(h) * 60 + int(m)


def clean(doc: dict) -> dict:
    return {k: v for k, v in doc.items() if k != "_id"}


def public_worker(w: dict) -> dict:
    out = clean(w)
    out.pop("pin_hash", None)
    out.pop("name_lower", None)
    return out


def period_bounds(month: str, period: int):
    year, mon = month.split("-")
    last = calendar.monthrange(int(year), int(mon))[1]
    if period == 1:
        return f"{month}-01", f"{month}-15"
    return f"{month}-16", f"{month}-{last:02d}"


def compute_day(rec: dict, daily_rate: float, overtime_rate: float) -> dict:
    """Hitung upah satu hari.

    Gaji pokok dihitung per HARI: hari penuh = tarif harian, setengah hari =
    setengah tarif harian. Tarif per jam (overtime_rate) HANYA dipakai untuk
    lembur (mulai 18.00) dan bonus tanpa istirahat siang (1,5 j) / sore (1 j).
    """
    day_type = "none"
    masuk = rec.get("masuk_at")
    pulang = rec.get("pulang_at")
    morning = afternoon = False
    if masuk:
        m_in = to_min(masuk)
        m_out = to_min(pulang) if pulang else None
        morning = m_in < MORNING_END
        afternoon = (m_in >= MORNING_END) or (m_out is None) or (m_out >= AFTERNOON_START)
        day_type = "full" if (morning and afternoon) else "half"
    base_pay = round(daily_rate if day_type == "full" else (daily_rate / 2 if day_type == "half" else 0))
    rest_hours = 0.0
    if rec.get("no_rest_siang") and morning and afternoon:
        rest_hours += LUNCH_REST_HOURS
    if rec.get("no_rest_sore") and afternoon:
        rest_hours += EVENING_REST_HOURS
    lembur_hours = float(rec.get("lembur_hours", 0) or 0) if rec.get("lembur") else 0.0
    rest_pay = round(rest_hours * overtime_rate)
    ot_pay = round(lembur_hours * overtime_rate)
    return {
        "day_type": day_type,
        "base_pay": base_pay,
        "rest_hours": round(rest_hours, 2),
        "rest_pay": rest_pay,
        "lembur_hours": lembur_hours,
        "overtime_pay": ot_pay,
        "total": base_pay + rest_pay + ot_pay,
    }


def day_status(rec: Optional[dict]) -> str:
    if not rec or not rec.get("masuk_at"):
        return "belum"
    if rec.get("pulang_at"):
        return "pulang"
    if rec.get("lembur"):
        return "lembur"
    events = rec.get("events", [])
    if events and events[-1].get("type") == "istirahat":
        return "istirahat"
    return "masuk"


# ---------- auth ----------

class AdminLogin(BaseModel):
    username: str = Field(min_length=1, max_length=128)
    password: str = Field(min_length=1, max_length=72)


class WorkerLogin(BaseModel):
    code: str = Field(min_length=1, max_length=64)
    pin: str = Field(pattern=r"^\d{4}$")


def create_token(user_id: str, role: str) -> str:
    now = datetime.now(timezone.utc)
    return jwt.encode(
        {"sub": user_id, "role": role, "iat": now, "exp": now + timedelta(days=JWT_DAYS)},
        JWT_SECRET,
        algorithm=JWT_ALG,
    )


def unauthorized():
    return HTTPException(status_code=401, detail="Sesi tidak valid, silakan masuk kembali")


async def current_user(credentials: Optional[HTTPAuthorizationCredentials] = Depends(bearer)):
    if not credentials or credentials.scheme.lower() != "bearer":
        raise unauthorized()
    try:
        claims = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=[JWT_ALG])
        sub, role = claims.get("sub"), claims.get("role")
        if not sub or role not in {"admin", "worker"}:
            raise unauthorized()
    except jwt.InvalidTokenError:
        raise unauthorized()
    if role == "admin":
        user = await db.admins.find_one({"id": sub}, {"_id": 0})
        if not user:
            raise unauthorized()
        return {"id": sub, "role": "admin", "username": user["username"]}
    worker = await db.workers.find_one({"id": sub, "active": True}, {"_id": 0})
    if not worker:
        raise unauthorized()
    return {"id": sub, "role": "worker", "worker": worker}


async def require_admin(user=Depends(current_user)):
    if user["role"] != "admin":
        raise HTTPException(status_code=403, detail="Khusus admin")
    return user


async def require_worker(user=Depends(current_user)):
    if user["role"] != "worker":
        raise HTTPException(status_code=403, detail="Khusus pekerja")
    return user


async def export_admin(token: str = Query(...)):
    try:
        claims = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
        if claims.get("role") != "admin":
            raise ValueError
    except Exception:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    return claims


@api_router.post("/auth/admin-login")
async def admin_login(body: AdminLogin):
    user = await db.admins.find_one({"username": body.username.strip()}, {"_id": 0})
    if not user or not pwd_ctx.verify(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Username atau kata sandi salah")
    return {"access_token": create_token(user["id"], "admin"), "token_type": "bearer", "role": "admin"}


@api_router.post("/auth/worker-login")
async def worker_login(body: WorkerLogin):
    code = body.code.strip().upper()
    worker = await db.workers.find_one({"code": code, "active": True}, {"_id": 0})
    if not worker or not pwd_ctx.verify(body.pin, worker["pin_hash"]):
        raise HTTPException(status_code=401, detail="Kode pekerja atau PIN salah")
    return {
        "access_token": create_token(worker["id"], "worker"),
        "token_type": "bearer",
        "role": "worker",
        "worker": public_worker(worker),
    }


@api_router.get("/auth/me")
async def auth_me(user=Depends(current_user)):
    if user["role"] == "admin":
        return {"role": "admin", "username": user["username"]}
    return {"role": "worker", "worker": public_worker(user["worker"])}


# ---------- workers (admin) ----------

class WorkerCreate(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    phone: str = Field(default="", max_length=30)
    group: str
    daily_rate: float = Field(gt=0)
    overtime_rate: float = Field(gt=0)
    pin: Optional[str] = Field(default=None, pattern=r"^\d{4}$")


class WorkerUpdate(BaseModel):
    name: Optional[str] = Field(default=None, min_length=2, max_length=80)
    phone: Optional[str] = Field(default=None, max_length=30)
    group: Optional[str] = None
    daily_rate: Optional[float] = Field(default=None, gt=0)
    overtime_rate: Optional[float] = Field(default=None, gt=0)
    active: Optional[bool] = None


@api_router.get("/groups")
async def get_groups(user=Depends(require_admin)):
    counts = await db.workers.aggregate([
        {"$match": {"active": True}},
        {"$group": {"_id": "$group", "count": {"$sum": 1}}},
    ]).to_list(50)
    count_map = {row["_id"]: row["count"] for row in counts}
    return [{"name": g, "count": count_map.get(g, 0)} for g in GROUPS]


@api_router.get("/workers")
async def get_workers(group: Optional[str] = None, include_inactive: bool = False, user=Depends(require_admin)):
    query = {} if include_inactive else {"active": True}
    if group and group != "Semua":
        query["group"] = group
    rows = await db.workers.find(query, {"_id": 0}).sort("name", 1).to_list(2000)
    return [public_worker(w) for w in rows]


async def next_worker_code() -> str:
    for _ in range(20):
        count = await db.workers.count_documents({})
        code = f"TG{count + 1:03d}"
        if not await db.workers.find_one({"code": code}):
            return code
    return f"TG{uuid.uuid4().hex[:5].upper()}"


@api_router.post("/workers", status_code=201)
async def create_worker(body: WorkerCreate, user=Depends(require_admin)):
    if body.group not in GROUPS:
        raise HTTPException(status_code=422, detail="Grup tidak dikenal")
    name_lower = body.name.strip().lower()
    existing = await db.workers.find_one({"name_lower": name_lower, "active": True}, {"_id": 0})
    if existing:
        raise HTTPException(status_code=409, detail=f"Pekerja bernama {body.name.strip()} sudah terdaftar ({existing['code']})")
    pin = body.pin or f"{random.randint(0, 9999):04d}"
    worker = {
        "id": str(uuid.uuid4()),
        "code": await next_worker_code(),
        "name": body.name.strip(),
        "name_lower": name_lower,
        "phone": body.phone.strip(),
        "group": body.group,
        "daily_rate": body.daily_rate,
        "overtime_rate": body.overtime_rate,
        "pin_hash": pwd_ctx.hash(pin),
        "active": True,
        "created_at": now_iso(),
    }
    await db.workers.insert_one(worker)
    return {**public_worker(worker), "pin": pin}


@api_router.put("/workers/{worker_id}")
async def update_worker(worker_id: str, body: WorkerUpdate, user=Depends(require_admin)):
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Pekerja tidak ditemukan")
    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    if "group" in updates and updates["group"] not in GROUPS:
        raise HTTPException(status_code=422, detail="Grup tidak dikenal")
    if "name" in updates:
        name_lower = updates["name"].strip().lower()
        dup = await db.workers.find_one({"name_lower": name_lower, "active": True, "id": {"$ne": worker_id}}, {"_id": 0})
        if dup:
            raise HTTPException(status_code=409, detail=f"Nama sudah dipakai pekerja lain ({dup['code']})")
        updates["name"] = updates["name"].strip()
        updates["name_lower"] = name_lower
    if updates:
        await db.workers.update_one({"id": worker_id}, {"$set": updates})
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    return public_worker(worker)


@api_router.post("/workers/{worker_id}/reset-pin")
async def reset_worker_pin(worker_id: str, user=Depends(require_admin)):
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Pekerja tidak ditemukan")
    pin = f"{random.randint(0, 9999):04d}"
    await db.workers.update_one({"id": worker_id}, {"$set": {"pin_hash": pwd_ctx.hash(pin)}})
    return {"id": worker_id, "code": worker["code"], "name": worker["name"], "pin": pin}


@api_router.delete("/workers/{worker_id}")
async def delete_worker(worker_id: str, user=Depends(require_admin)):
    worker = await db.workers.find_one({"id": worker_id}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Pekerja tidak ditemukan")
    has_records = await db.attendance.count_documents({"worker_id": worker_id})
    if has_records:
        await db.workers.update_one({"id": worker_id}, {"$set": {"active": False}})
        return {"id": worker_id, "deleted": False, "deactivated": True}
    await db.workers.delete_one({"id": worker_id})
    return {"id": worker_id, "deleted": True, "deactivated": False}


# ---------- attendance (admin) ----------

class TapInput(BaseModel):
    worker_id: str
    type: Literal["masuk", "istirahat", "lembur", "pulang"]
    date: Optional[str] = Field(default=None, pattern=r"^\d{4}-\d{2}-\d{2}$")
    time: Optional[str] = Field(default=None, pattern=r"^\d{2}:\d{2}$")


class AttendanceEdit(BaseModel):
    masuk_at: Optional[str] = Field(default=None, pattern=r"^\d{2}:\d{2}$")
    pulang_at: Optional[str] = Field(default=None, pattern=r"^\d{2}:\d{2}$")
    no_rest_siang: bool = False
    no_rest_sore: bool = False
    lembur: bool = False
    lembur_hours: Optional[float] = Field(default=None, ge=0, le=12)
    lembur_note: str = Field(default="", max_length=240)


def attendance_row(worker: dict, rec: Optional[dict]) -> dict:
    return {
        "worker": public_worker(worker),
        "record": clean(rec) if rec else None,
        "status": day_status(rec),
        "calc": compute_day(rec, worker["daily_rate"], worker["overtime_rate"]) if rec and rec.get("masuk_at") else None,
    }


@api_router.get("/attendance")
async def get_attendance(
    date: str = Query(default=None),
    group: Optional[str] = None,
    user=Depends(require_admin),
):
    day = date or today_wib()
    query = {"active": True}
    if group and group != "Semua":
        query["group"] = group
    workers = await db.workers.find(query, {"_id": 0}).sort("name", 1).to_list(2000)
    recs = await db.attendance.find({"date": day}, {"_id": 0}).to_list(5000)
    rec_map = {r["worker_id"]: r for r in recs}
    return {"date": day, "rows": [attendance_row(w, rec_map.get(w["id"])) for w in workers]}


@api_router.post("/attendance/tap")
async def tap_attendance(body: TapInput, user=Depends(require_admin)):
    worker = await db.workers.find_one({"id": body.worker_id, "active": True}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Pekerja tidak ditemukan")
    day = body.date or today_wib()
    t = body.time or now_hhmm()
    rec = await db.attendance.find_one({"worker_id": body.worker_id, "date": day}, {"_id": 0})
    if body.type == "masuk":
        if rec and rec.get("masuk_at"):
            raise HTTPException(status_code=400, detail="Pekerja sudah tercatat masuk")
    else:
        if not rec or not rec.get("masuk_at"):
            raise HTTPException(status_code=400, detail="Catat jam masuk dulu")
        if rec.get("pulang_at"):
            raise HTTPException(status_code=400, detail="Pekerja sudah pulang")
    if not rec:
        rec = {
            "id": str(uuid.uuid4()),
            "worker_id": body.worker_id,
            "date": day,
            "masuk_at": None,
            "pulang_at": None,
            "no_rest_siang": False,
            "no_rest_sore": False,
            "lembur": False,
            "lembur_hours": 0.0,
            "lembur_note": "",
            "events": [],
        }
    if body.type == "masuk":
        rec["masuk_at"] = t
    elif body.type == "lembur":
        rec["lembur"] = True
    elif body.type == "pulang":
        rec["pulang_at"] = t
        if rec.get("lembur") and not rec.get("lembur_manual"):
            rec["lembur_hours"] = round(max(to_min(t) - EVENING_REST_END, 0) / 30) / 2
    rec["events"].append({"type": body.type, "at": f"{day}T{t}:00+07:00"})
    rec["updated_at"] = now_iso()
    await db.attendance.replace_one({"id": rec["id"]}, rec, upsert=True)
    return attendance_row(worker, rec)


@api_router.put("/attendance/{rec_id}")
async def edit_attendance(rec_id: str, body: AttendanceEdit, user=Depends(require_admin)):
    rec = await db.attendance.find_one({"id": rec_id}, {"_id": 0})
    if not rec:
        raise HTTPException(status_code=404, detail="Catatan absensi tidak ditemukan")
    worker = await db.workers.find_one({"id": rec["worker_id"]}, {"_id": 0})
    if not worker:
        raise HTTPException(status_code=404, detail="Pekerja tidak ditemukan")
    rec["masuk_at"] = body.masuk_at or None
    rec["pulang_at"] = body.pulang_at or None
    rec["no_rest_siang"] = body.no_rest_siang
    rec["no_rest_sore"] = body.no_rest_sore
    rec["lembur"] = body.lembur
    rec["lembur_note"] = body.lembur_note
    if body.lembur_hours is not None:
        rec["lembur_hours"] = body.lembur_hours
        rec["lembur_manual"] = True
    elif body.lembur and rec.get("pulang_at"):
        rec["lembur_hours"] = round(max(to_min(rec["pulang_at"]) - EVENING_REST_END, 0) / 30) / 2
        rec["lembur_manual"] = False
    elif not body.lembur:
        rec["lembur_hours"] = 0.0
        rec["lembur_manual"] = False
    rec["updated_at"] = now_iso()
    await db.attendance.replace_one({"id": rec_id}, rec)
    return attendance_row(worker, rec)


@api_router.delete("/attendance/{rec_id}")
async def delete_attendance(rec_id: str, user=Depends(require_admin)):
    result = await db.attendance.delete_one({"id": rec_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Catatan absensi tidak ditemukan")
    return {"deleted": True, "id": rec_id}


# ---------- dashboard (admin) ----------

@api_router.get("/dashboard")
async def get_dashboard(date: str = Query(default=None), user=Depends(require_admin)):
    day = date or today_wib()
    workers = await db.workers.find({"active": True}, {"_id": 0}).to_list(2000)
    recs = await db.attendance.find({"date": day}, {"_id": 0}).to_list(5000)
    rec_map = {r["worker_id"]: r for r in recs}
    wmap = {w["id"]: w for w in workers}
    present = sum(1 for r in rec_map.values() if r.get("masuk_at") and r["worker_id"] in wmap)
    pulang_done = sum(1 for r in rec_map.values() if r.get("pulang_at") and r["worker_id"] in wmap)
    lembur = sum(1 for r in rec_map.values() if r.get("lembur") and r["worker_id"] in wmap)
    est = 0
    for wid, r in rec_map.items():
        w = wmap.get(wid)
        if w and r.get("masuk_at"):
            est += compute_day(r, w["daily_rate"], w["overtime_rate"])["total"]
    groups = []
    for g in GROUPS:
        members = [w for w in workers if w["group"] == g]
        hadir = sum(1 for w in members if rec_map.get(w["id"], {}).get("masuk_at"))
        groups.append({"group": g, "total": len(members), "present": hadir})
    return {
        "date": day,
        "active_workers": len(workers),
        "present": present,
        "pulang_done": pulang_done,
        "lembur": lembur,
        "absent": max(len(workers) - present, 0),
        "est_pay_today": est,
        "groups": groups,
    }


# ---------- payroll & payments (admin) ----------

class PaymentItem(BaseModel):
    worker_id: str
    amount: float = Field(ge=0)


class PaymentCreate(BaseModel):
    month: str = Field(pattern=r"^\d{4}-\d{2}$")
    period: Literal[1, 2]
    method: Literal["mandor", "langsung"]
    mandor_name: str = Field(default="", max_length=80)
    note: str = Field(default="", max_length=240)
    items: List[PaymentItem]


def payroll_row(worker: dict, recs: list, paid: float) -> dict:
    days = half_days = 0
    rest = ot = base = rest_pay = ot_pay = 0.0
    for r in recs:
        if not r.get("masuk_at"):
            continue
        c = compute_day(r, worker["daily_rate"], worker["overtime_rate"])
        if c["day_type"] == "full":
            days += 1
        elif c["day_type"] == "half":
            half_days += 1
        rest += c["rest_hours"]
        ot += c["lembur_hours"]
        base += c["base_pay"]
        rest_pay += c["rest_pay"]
        ot_pay += c["overtime_pay"]
    gross = base + rest_pay + ot_pay
    return {
        "worker_id": worker["id"],
        "code": worker["code"],
        "name": worker["name"],
        "group": worker["group"],
        "daily_rate": worker["daily_rate"],
        "overtime_rate": worker["overtime_rate"],
        "days": days,
        "half_days": half_days,
        "rest_hours": round(rest, 2),
        "overtime_hours": round(ot, 2),
        "base_pay": round(base),
        "rest_pay": round(rest_pay),
        "overtime_pay": round(ot_pay),
        "gross": round(gross),
        "paid": round(paid),
        "remaining": round(max(gross - paid, 0)),
    }


@api_router.get("/payroll")
async def get_payroll(
    month: str = Query(..., pattern=r"^\d{4}-\d{2}$"),
    period: int = Query(default=1, ge=1, le=2),
    group: Optional[str] = None,
    user=Depends(require_admin),
):
    start, end = period_bounds(month, period)
    wq = {"active": True}
    if group and group != "Semua":
        wq["group"] = group
    workers = await db.workers.find(wq, {"_id": 0}).sort("name", 1).to_list(2000)
    wids = {w["id"] for w in workers}
    recs = await db.attendance.find({"date": {"$gte": start, "$lte": end}, "worker_id": {"$in": list(wids)}}, {"_id": 0}).to_list(20000)
    pays = await db.payments.find({"month": month, "period": period}, {"_id": 0}).to_list(1000)
    paid_map = {}
    for p in pays:
        for s in p["splits"]:
            paid_map[s["worker_id"]] = paid_map.get(s["worker_id"], 0) + s["amount"]
    rows = []
    for w in workers:
        worker_recs = [r for r in recs if r["worker_id"] == w["id"]]
        rows.append(payroll_row(w, worker_recs, paid_map.get(w["id"], 0)))
    return {
        "month": month,
        "period": period,
        "start_date": start,
        "end_date": end,
        "rows": rows,
        "total_gross": sum(r["gross"] for r in rows),
        "total_paid": sum(r["paid"] for r in rows),
        "total_remaining": sum(r["remaining"] for r in rows),
    }


@api_router.get("/payments")
async def get_payments(month: str = Query(...), period: int = Query(default=1, ge=1, le=2), user=Depends(require_admin)):
    rows = await db.payments.find({"month": month, "period": period}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return rows


@api_router.post("/payments", status_code=201)
async def create_payment(body: PaymentCreate, user=Depends(require_admin)):
    if not body.items:
        raise HTTPException(status_code=422, detail="Pilih minimal satu pekerja")
    if body.method == "mandor" and not body.mandor_name.strip():
        raise HTTPException(status_code=422, detail="Nama mandor wajib diisi")
    splits = []
    for item in body.items:
        worker = await db.workers.find_one({"id": item.worker_id}, {"_id": 0})
        if not worker:
            raise HTTPException(status_code=404, detail="Pekerja tidak ditemukan")
        if item.amount <= 0:
            raise HTTPException(status_code=422, detail=f"Jumlah untuk {worker['name']} harus lebih dari 0")
        splits.append({"worker_id": worker["id"], "worker_name": worker["name"], "amount": round(item.amount)})
    payment = {
        "id": str(uuid.uuid4()),
        "month": body.month,
        "period": body.period,
        "method": body.method,
        "mandor_name": body.mandor_name.strip(),
        "note": body.note.strip(),
        "splits": splits,
        "total": sum(s["amount"] for s in splits),
        "created_by": user["username"],
        "created_at": now_iso(),
    }
    await db.payments.insert_one(payment)
    return clean(payment)


@api_router.delete("/payments/{payment_id}")
async def delete_payment(payment_id: str, user=Depends(require_admin)):
    result = await db.payments.delete_one({"id": payment_id})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Pembayaran tidak ditemukan")
    return {"deleted": True, "id": payment_id}


# ---------- worker self-service ----------

@api_router.get("/worker/me")
async def worker_me(user=Depends(require_worker)):
    return public_worker(user["worker"])


@api_router.get("/worker/attendance")
async def worker_attendance(month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), user=Depends(require_worker)):
    worker = user["worker"]
    recs = await db.attendance.find(
        {"worker_id": worker["id"], "date": {"$gte": f"{month}-01", "$lte": f"{month}-31"}},
        {"_id": 0},
    ).sort("date", -1).to_list(100)
    rows = []
    for r in recs:
        rows.append({
            "record": clean(r),
            "status": day_status(r),
            "calc": compute_day(r, worker["daily_rate"], worker["overtime_rate"]) if r.get("masuk_at") else None,
        })
    return {"month": month, "rows": rows}


@api_router.get("/worker/payroll")
async def worker_payroll(month: str = Query(...), period: int = Query(default=1, ge=1, le=2), user=Depends(require_worker)):
    worker = user["worker"]
    start, end = period_bounds(month, period)
    recs = await db.attendance.find({"worker_id": worker["id"], "date": {"$gte": start, "$lte": end}}, {"_id": 0}).to_list(100)
    row = payroll_row(worker, recs, 0)
    pays = await db.payments.find({"month": month, "period": period, "splits.worker_id": worker["id"]}, {"_id": 0}).sort("created_at", -1).to_list(100)
    received = []
    paid = 0
    for p in pays:
        for s in p["splits"]:
            if s["worker_id"] == worker["id"]:
                paid += s["amount"]
                received.append({
                    "id": p["id"],
                    "method": p["method"],
                    "mandor_name": p["mandor_name"],
                    "note": p["note"],
                    "amount": s["amount"],
                    "created_at": p["created_at"],
                })
    return {
        "month": month,
        "period": period,
        "start_date": start,
        "end_date": end,
        "days": row["days"],
        "half_days": row["half_days"],
        "rest_hours": row["rest_hours"],
        "overtime_hours": row["overtime_hours"],
        "base_pay": row["base_pay"],
        "rest_pay": row["rest_pay"],
        "overtime_pay": row["overtime_pay"],
        "gross": row["gross"],
        "paid": round(paid),
        "remaining": round(max(row["gross"] - paid, 0)),
        "payments": received,
    }


# ---------- Excel export (admin) ----------

HEADER_FILL = PatternFill("solid", fgColor="E05A36")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def style_sheet(ws, headers, widths):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = widths[col - 1]


def xlsx_response(wb: Workbook, filename: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@api_router.get("/export/payroll")
async def export_payroll(month: str, period: int = 1, group: Optional[str] = None, user=Depends(export_admin)):
    start, end = period_bounds(month, period)
    wq = {"active": True}
    if group and group != "Semua":
        wq["group"] = group
    workers = await db.workers.find(wq, {"_id": 0}).sort([("group", 1), ("name", 1)]).to_list(2000)
    wids = {w["id"] for w in workers}
    recs = await db.attendance.find({"date": {"$gte": start, "$lte": end}, "worker_id": {"$in": list(wids)}}, {"_id": 0}).to_list(20000)
    pays = await db.payments.find({"month": month, "period": period}, {"_id": 0}).to_list(1000)
    paid_map = {}
    for p in pays:
        for s in p["splits"]:
            paid_map[s["worker_id"]] = paid_map.get(s["worker_id"], 0) + s["amount"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Gaji"
    headers = ["Kode", "Nama", "Grup", "Hari Penuh", "Setengah Hari", "Jam No-Rest", "Jam Lembur", "Upah Harian (Rp)", "Upah No-Rest (Rp)", "Upah Lembur (Rp)", "Total (Rp)", "Terbayar (Rp)", "Sisa (Rp)"]
    style_sheet(ws, headers, [10, 26, 12, 10, 12, 11, 11, 16, 16, 16, 16, 16, 16])
    for i, w in enumerate(workers, start=2):
        row = payroll_row(w, [r for r in recs if r["worker_id"] == w["id"]], paid_map.get(w["id"], 0))
        vals = [row["code"], row["name"], row["group"], row["days"], row["half_days"], row["rest_hours"], row["overtime_hours"], row["base_pay"], row["rest_pay"], row["overtime_pay"], row["gross"], row["paid"], row["remaining"]]
        for col, val in enumerate(vals, start=1):
            ws.cell(row=i, column=col, value=val)
    return xlsx_response(wb, f"gaji-{month}-periode{period}.xlsx")


@api_router.get("/export/attendance")
async def export_attendance(month: str, period: int = 1, group: Optional[str] = None, user=Depends(export_admin)):
    start, end = period_bounds(month, period)
    wq = {"active": True}
    if group and group != "Semua":
        wq["group"] = group
    workers = await db.workers.find(wq, {"_id": 0}).to_list(2000)
    wmap = {w["id"]: w for w in workers}
    recs = await db.attendance.find({"date": {"$gte": start, "$lte": end}, "worker_id": {"$in": list(wmap)}}, {"_id": 0}).sort([("date", 1), ("worker_id", 1)]).to_list(20000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Absensi"
    headers = ["Tanggal", "Kode", "Nama", "Grup", "Masuk", "Pulang", "Tipe Hari", "Tanpa Istirahat Siang", "Tanpa Istirahat Sore", "Jam Lembur", "Catatan Lembur", "Upah Hari (Rp)"]
    style_sheet(ws, headers, [12, 10, 26, 12, 8, 8, 11, 12, 12, 11, 24, 16])
    i = 2
    for r in recs:
        w = wmap.get(r["worker_id"])
        if not w:
            continue
        calc = compute_day(r, w["daily_rate"], w["overtime_rate"]) if r.get("masuk_at") else None
        day_label = "-" if not calc else ("Penuh" if calc["day_type"] == "full" else "Setengah")
        vals = [
            r["date"], w["code"], w["name"], w["group"], r.get("masuk_at") or "-", r.get("pulang_at") or "-",
            day_label,
            "Ya" if r.get("no_rest_siang") else "Tidak", "Ya" if r.get("no_rest_sore") else "Tidak",
            calc["lembur_hours"] if calc else 0,
            r.get("lembur_note") or "", calc["total"] if calc else 0,
        ]
        for col, val in enumerate(vals, start=1):
            ws.cell(row=i, column=col, value=val)
        i += 1
    return xlsx_response(wb, f"absensi-{month}-periode{period}.xlsx")


@api_router.get("/")
async def root():
    return {"message": "TukangGaji Pro API v2 (gaji per hari)"}


# ---------- app setup ----------

@asynccontextmanager
async def lifespan(app: FastAPI):
    await db.admins.update_one(
        {"username": ADMIN_USERNAME},
        {"$setOnInsert": {"id": str(uuid.uuid4()), "username": ADMIN_USERNAME, "password_hash": pwd_ctx.hash(ADMIN_PASSWORD), "created_at": now_iso()}},
        upsert=True,
    )
    await db.workers.create_index("code", unique=True)
    await db.attendance.create_index([("worker_id", 1), ("date", 1)], unique=True)
    yield


app = FastAPI(lifespan=lifespan)
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)
