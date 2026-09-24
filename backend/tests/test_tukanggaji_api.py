"""TukangGaji Pro backend regression tests - model GAJI PER HARI (iteration 2).

Covers: health, auth (admin+worker PIN 8366), workers CRUD dengan daily_rate,
perhitungan gaji per-hari (hari penuh/setengah/no-rest/lembur), payroll periode 2
Sep 2026, payments mandor/langsung, role guards, worker self-service, ekspor Excel
dengan kolom baru (Hari Penuh, Setengah Hari, Jam No-Rest, Jam Lembur, Tipe Hari).

Data seeded yang diverifikasi (JANGAN diubah): TG001 absensi 2026-09-24
07:00-20:30 + no-rest siang&sore + lembur 2,5j -> total Rp260.000;
pembayaran Rp100.000 via mandor Pak RT periode 2.
"""
import io

import pytest
import requests
from openpyxl import load_workbook

MONTH = "2026-09"  # current seeded month (today = 2026-09-24 WIB)
TODAY = "2026-09-24"


# ---------- health ----------

class TestHealth:
    def test_root(self, api_client, base_url):
        resp = api_client.get(f"{base_url}/api/")
        assert resp.status_code == 200
        assert "TukangGaji" in resp.json()["message"]


# ---------- auth ----------

class TestAuth:
    def test_admin_login_success(self, api_client, base_url):
        resp = api_client.post(f"{base_url}/api/auth/admin-login", json={"username": "admin", "password": "tukangpro123"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["role"] == "admin"
        assert data["access_token"]

    def test_admin_login_wrong_password(self, api_client, base_url):
        resp = api_client.post(f"{base_url}/api/auth/admin-login", json={"username": "admin", "password": "salah123"})
        assert resp.status_code == 401

    def test_worker_login_success_pin_8366(self, api_client, base_url):
        resp = api_client.post(f"{base_url}/api/auth/worker-login", json={"code": "TG001", "pin": "8366"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["role"] == "worker"
        assert data["worker"]["code"] == "TG001"
        assert data["worker"]["name"] == "Budi Santoso"
        assert data["worker"]["daily_rate"] == 160000
        assert data["worker"]["overtime_rate"] == 20000
        assert "pin_hash" not in data["worker"]
        assert "hourly_rate" not in data["worker"]  # old field must be gone

    def test_worker_login_wrong_pin(self, api_client, base_url):
        resp = api_client.post(f"{base_url}/api/auth/worker-login", json={"code": "TG001", "pin": "0000"})
        assert resp.status_code == 401

    def test_worker_cannot_access_admin_endpoint(self, api_client, base_url, worker_token):
        resp = api_client.get(f"{base_url}/api/workers", headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 403

    def test_admin_cannot_access_worker_endpoint(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/worker/me")
        assert resp.status_code == 403

    def test_no_token_unauthorized(self, base_url):
        resp = requests.get(f"{base_url}/api/dashboard")
        assert resp.status_code == 401


# ---------- workers CRUD (daily_rate model) ----------

class TestWorkers:
    def test_create_worker_daily_rate_returns_code_and_pin(self, admin_client, base_url, test_worker):
        assert test_worker["code"].startswith("TG")
        assert test_worker["pin"] == "4321"
        assert test_worker["group"] == "Kuli"
        assert test_worker["daily_rate"] == 160000
        assert test_worker["overtime_rate"] == 20000
        # persistence check
        get = admin_client.get(f"{base_url}/api/workers")
        persisted = [w for w in get.json() if w["id"] == test_worker["id"]][0]
        assert persisted["daily_rate"] == 160000

    def test_duplicate_name_rejected_409(self, admin_client, base_url, test_worker):
        resp = admin_client.post(f"{base_url}/api/workers", json={
            "name": "test pekerja pytest", "phone": "", "group": "Tukang",
            "daily_rate": 100000, "overtime_rate": 10000,
        })
        assert resp.status_code == 409
        assert "sudah terdaftar" in resp.json()["detail"].lower()

    def test_update_worker_daily_rate(self, admin_client, base_url, test_worker):
        resp = admin_client.put(f"{base_url}/api/workers/{test_worker['id']}", json={"daily_rate": 170000, "overtime_rate": 22000})
        assert resp.status_code == 200
        data = resp.json()
        assert data["daily_rate"] == 170000
        assert data["overtime_rate"] == 22000
        # GET verify persistence
        get = admin_client.get(f"{base_url}/api/workers")
        persisted = [w for w in get.json() if w["id"] == test_worker["id"]][0]
        assert persisted["daily_rate"] == 170000
        # restore
        admin_client.put(f"{base_url}/api/workers/{test_worker['id']}", json={"daily_rate": 160000, "overtime_rate": 20000})

    def test_reset_pin_and_login(self, api_client, admin_client, base_url, test_worker):
        resp = admin_client.post(f"{base_url}/api/workers/{test_worker['id']}/reset-pin")
        assert resp.status_code == 200
        new_pin = resp.json()["pin"]
        assert len(new_pin) == 4
        login = api_client.post(f"{base_url}/api/auth/worker-login", json={"code": test_worker["code"], "pin": new_pin})
        assert login.status_code == 200

    def test_invalid_group_rejected(self, admin_client, base_url):
        resp = admin_client.post(f"{base_url}/api/workers", json={
            "name": "TEST Grup Ngawur", "phone": "", "group": "Sihir",
            "daily_rate": 100000, "overtime_rate": 10000,
        })
        assert resp.status_code == 422

    def test_missing_daily_rate_rejected(self, admin_client, base_url):
        resp = admin_client.post(f"{base_url}/api/workers", json={
            "name": "TEST Tanpa Tarif", "phone": "", "group": "Kuli", "overtime_rate": 10000,
        })
        assert resp.status_code == 422


# ---------- perhitungan gaji PER HARI ----------

def _create_day(admin_client, base_url, wid, date, masuk, pulang, no_siang=False, no_sore=False, lembur=False, lembur_hours=None, note=""):
    """Tap masuk+pulang lalu edit via PUT untuk kontrol penuh. Mengembalikan calc."""
    r = admin_client.post(f"{base_url}/api/attendance/tap", json={"worker_id": wid, "type": "masuk", "date": date, "time": masuk})
    assert r.status_code == 200, r.text
    r = admin_client.post(f"{base_url}/api/attendance/tap", json={"worker_id": wid, "type": "pulang", "date": date, "time": pulang})
    assert r.status_code == 200, r.text
    rec_id = r.json()["record"]["id"]
    edit = admin_client.put(f"{base_url}/api/attendance/{rec_id}", json={
        "masuk_at": masuk, "pulang_at": pulang,
        "no_rest_siang": no_siang, "no_rest_sore": no_sore,
        "lembur": lembur, "lembur_hours": lembur_hours, "lembur_note": note,
    })
    assert edit.status_code == 200, edit.text
    return rec_id, edit.json()["calc"]


class TestDailyWageMath:
    """Skenario pada TEST worker (harian 160.000, lembur 20.000/jam)."""

    def test_full_day_with_norest_and_lembur(self, admin_client, base_url, test_worker):
        # 07:00-20:30 + no-rest siang&sore + lembur 2,5j -> 160rb + 50rb + 50rb = 260rb
        rec_id, calc = _create_day(admin_client, base_url, test_worker["id"], "2026-09-20",
                                   "07:00", "20:30", no_siang=True, no_sore=True,
                                   lembur=True, lembur_hours=2.5, note="TEST pasang rangka")
        assert calc["day_type"] == "full"
        assert calc["base_pay"] == 160000
        assert calc["rest_hours"] == 2.5
        assert calc["rest_pay"] == 50000
        assert calc["lembur_hours"] == 2.5
        assert calc["overtime_pay"] == 50000
        assert calc["total"] == 260000

    def test_half_day_morning(self, admin_client, base_url, test_worker):
        # 07:00-11:45 -> setengah hari -> 80.000
        rec_id, calc = _create_day(admin_client, base_url, test_worker["id"], "2026-09-21", "07:00", "11:45")
        assert calc["day_type"] == "half"
        assert calc["base_pay"] == 80000
        assert calc["total"] == 80000

    def test_half_day_afternoon(self, admin_client, base_url, test_worker):
        # 13:00-17:00 -> setengah hari -> 80.000
        rec_id, calc = _create_day(admin_client, base_url, test_worker["id"], "2026-09-22", "13:00", "17:00")
        assert calc["day_type"] == "half"
        assert calc["base_pay"] == 80000
        assert calc["total"] == 80000

    def test_full_day_plain(self, admin_client, base_url, test_worker):
        # 07:00-17:00 polos -> hari penuh -> 160.000, tanpa no-rest/lembur
        rec_id, calc = _create_day(admin_client, base_url, test_worker["id"], "2026-09-23", "07:00", "17:00")
        assert calc["day_type"] == "full"
        assert calc["base_pay"] == 160000
        assert calc["rest_hours"] == 0
        assert calc["rest_pay"] == 0
        assert calc["overtime_pay"] == 0
        assert calc["total"] == 160000

    def test_seeded_tg001_today_calc(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/attendance?date={TODAY}")
        assert resp.status_code == 200
        row = [r for r in resp.json()["rows"] if r["worker"]["code"] == "TG001"][0]
        assert row["record"]["masuk_at"] == "07:00"
        assert row["record"]["pulang_at"] == "20:30"
        assert row["record"]["no_rest_siang"] is True
        assert row["record"]["no_rest_sore"] is True
        calc = row["calc"]
        assert calc["day_type"] == "full"
        assert calc["base_pay"] == 160000
        assert calc["rest_pay"] == 50000
        assert calc["overtime_pay"] == 50000
        assert calc["total"] == 260000


# ---------- dashboard ----------

class TestDashboard:
    def test_dashboard_today(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/dashboard")
        assert resp.status_code == 200
        data = resp.json()
        assert data["date"] == TODAY
        assert data["active_workers"] >= 1
        assert data["present"] >= 1
        assert data["est_pay_today"] >= 260000
        tukang = [g for g in data["groups"] if g["group"] == "Tukang"][0]
        assert tukang["total"] >= 1 and tukang["present"] >= 1


# ---------- payroll & payments ----------

class TestPayroll:
    def test_payroll_period2_tg001_fields_and_totals(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2")
        assert resp.status_code == 200
        data = resp.json()
        assert data["start_date"] == "2026-09-16" and data["end_date"] == "2026-09-30"
        row = [r for r in data["rows"] if r["code"] == "TG001"][0]
        # field rows harus memuat breakdown per-hari
        assert row["days"] == 1
        assert row["half_days"] == 0
        assert row["rest_hours"] == 2.5
        assert row["overtime_hours"] == 2.5
        assert row["base_pay"] == 160000
        assert row["rest_pay"] == 50000
        assert row["overtime_pay"] == 50000
        assert row["gross"] == 260000
        assert row["paid"] == 100000
        assert row["remaining"] == 160000
        assert row["daily_rate"] == 160000
        assert "hourly_rate" not in row

    def test_payment_history_contains_mandor(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/payments?month={MONTH}&period=2")
        assert resp.status_code == 200
        payments = resp.json()
        assert any(p["method"] == "mandor" and p["mandor_name"] == "Pak RT" and p["total"] == 100000 for p in payments)

    def test_mandor_payment_requires_name(self, admin_client, base_url, test_worker):
        resp = admin_client.post(f"{base_url}/api/payments", json={
            "month": MONTH, "period": 2, "method": "mandor", "mandor_name": "", "note": "",
            "items": [{"worker_id": test_worker["id"], "amount": 5000}],
        })
        assert resp.status_code == 422

    def test_create_payment_reduces_remaining_then_delete(self, admin_client, base_url, test_worker):
        # beri TEST worker 1 hari penuh di periode 2 supaya gross > 0
        wid = test_worker["id"]
        rows = admin_client.get(f"{base_url}/api/attendance?date=2026-09-23").json()["rows"]
        rec = [r for r in rows if r["worker"]["id"] == wid][0]["record"]
        if not rec:
            _create_day(admin_client, base_url, wid, "2026-09-23", "07:00", "17:00")
        before = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        brow = [r for r in before["rows"] if r["code"] == test_worker["code"]][0]

        resp = admin_client.post(f"{base_url}/api/payments", json={
            "month": MONTH, "period": 2, "method": "langsung", "mandor_name": "", "note": "TEST bayar",
            "items": [{"worker_id": wid, "amount": 5000}],
        })
        assert resp.status_code == 201
        pid = resp.json()["id"]

        after = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        arow = [r for r in after["rows"] if r["code"] == test_worker["code"]][0]
        assert arow["paid"] == brow["paid"] + 5000
        assert arow["remaining"] == max(brow["gross"] - arow["paid"], 0)

        dele = admin_client.delete(f"{base_url}/api/payments/{pid}")
        assert dele.status_code == 200
        after2 = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        arow2 = [r for r in after2["rows"] if r["code"] == test_worker["code"]][0]
        assert arow2["paid"] == brow["paid"]


# ---------- worker self-service ----------

class TestWorkerSelfService:
    def test_worker_payroll_period2(self, api_client, base_url, worker_token):
        resp = api_client.get(f"{base_url}/api/worker/payroll?month={MONTH}&period=2",
                              headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["days"] == 1 and data["half_days"] == 0
        assert data["rest_hours"] == 2.5 and data["overtime_hours"] == 2.5
        assert data["base_pay"] == 160000
        assert data["rest_pay"] == 50000
        assert data["overtime_pay"] == 50000
        assert data["gross"] == 260000
        assert data["paid"] == 100000
        assert data["remaining"] == 160000
        assert any(p["method"] == "mandor" and p["mandor_name"] == "Pak RT" and p["amount"] == 100000 for p in data["payments"])

    def test_worker_attendance_month(self, api_client, base_url, worker_token):
        resp = api_client.get(f"{base_url}/api/worker/attendance?month={MONTH}",
                              headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 200
        rows = resp.json()["rows"]
        today_row = [r for r in rows if r["record"]["date"] == TODAY]
        assert today_row
        calc = today_row[0]["calc"]
        assert calc["day_type"] == "full"
        assert calc["total"] == 260000


# ---------- Excel export ----------

class TestExport:
    def _load(self, resp):
        assert resp.status_code == 200
        assert resp.headers["Content-Type"].startswith("application/vnd.openxmlformats")
        assert resp.content[:2] == b"PK"
        return load_workbook(io.BytesIO(resp.content))

    def test_export_payroll_columns_new_model(self, base_url, admin_token):
        resp = requests.get(f"{base_url}/api/export/payroll?month={MONTH}&period=2&token={admin_token}")
        wb = self._load(resp)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        for col in ["Hari Penuh", "Setengah Hari", "Jam No-Rest", "Jam Lembur",
                    "Upah Harian (Rp)", "Upah No-Rest (Rp)", "Upah Lembur (Rp)"]:
            assert col in headers, f"Kolom '{col}' tidak ada: {headers}"
        # baris TG001: 1 hari penuh, 2,5 jam no-rest, 2,5 jam lembur, total 260.000
        kode_idx = headers.index("Kode") + 1
        tg_row = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=kode_idx).value == "TG001":
                tg_row = r
                break
        assert tg_row, "TG001 tidak ada di ekspor gaji"
        assert ws.cell(row=tg_row, column=headers.index("Hari Penuh") + 1).value == 1
        assert ws.cell(row=tg_row, column=headers.index("Jam No-Rest") + 1).value == 2.5
        assert ws.cell(row=tg_row, column=headers.index("Jam Lembur") + 1).value == 2.5
        assert ws.cell(row=tg_row, column=headers.index("Upah Harian (Rp)") + 1).value == 160000
        assert ws.cell(row=tg_row, column=headers.index("Total (Rp)") + 1).value == 260000
        assert ws.cell(row=tg_row, column=headers.index("Terbayar (Rp)") + 1).value == 100000
        assert ws.cell(row=tg_row, column=headers.index("Sisa (Rp)") + 1).value == 160000

    def test_export_attendance_tipe_hari_column(self, base_url, admin_token):
        resp = requests.get(f"{base_url}/api/export/attendance?month={MONTH}&period=2&token={admin_token}")
        wb = self._load(resp)
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Tipe Hari" in headers, f"Kolom 'Tipe Hari' tidak ada: {headers}"
        tipe_idx = headers.index("Tipe Hari") + 1
        kode_idx = headers.index("Kode") + 1
        tgl_idx = headers.index("Tanggal") + 1
        tg_row = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=kode_idx).value == "TG001" and ws.cell(row=r, column=tgl_idx).value == TODAY:
                tg_row = r
                break
        assert tg_row, "Baris TG001 hari ini tidak ada di ekspor absensi"
        assert ws.cell(row=tg_row, column=tipe_idx).value == "Penuh"

    def test_export_requires_admin_token(self, base_url, worker_token):
        resp = requests.get(f"{base_url}/api/export/payroll?month={MONTH}&period=2&token={worker_token}")
        assert resp.status_code == 401
