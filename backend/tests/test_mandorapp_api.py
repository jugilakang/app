"""MandorApp backend regression tests - iteration 3.

Perubahan sesi ini yang dites:
1. NO-REST SIANG = 1 JAM (11.30-12.30), BUKAN 1,5 jam; no-rest sore tetap 1 jam.
   Keduanya x tarif lembur. Gaji pokok tetap per hari (penuh / setengah).
2. Payroll memuat total_days & total_overtime_hours (GRAND TOTAL periode) yang
   konsisten dengan jumlah baris per pekerja.
3. Ekspor Excel gaji memuat baris GRAND TOTAL di baris terakhir (jumlah hari,
   jam lembur, total rupiah).
4. Rebrand: tidak ada teks 'TukangGaji' pada respons API root (dilaporkan, lihat
   test_root_rebrand_note).

DATA ASLI user (JANGAN diubah): Bayu TG001, Reza TG002, Obeng TG003.
Semua tes memakai pekerja TEST_ yang dibuat & dihapus oleh fixture.
Periode uji: 2026-09 periode 2 (hari ini 24 Sep 2026 WIB).
"""
import io

import pytest
import requests
from openpyxl import load_workbook

from conftest import ensure_day

MONTH = "2026-09"  # bulan berjalan (today = 2026-09-24 WIB)


# ---------- health / rebrand ----------

class TestHealth:
    def test_root(self, api_client, base_url):
        resp = api_client.get(f"{base_url}/api/")
        assert resp.status_code == 200
        assert resp.json()["message"]

    def test_root_rebrand_note(self, api_client, base_url):
        # Rebrand check: API root message seharusnya tidak lagi menyebut 'TukangGaji'.
        # Dicatat sebagai temuan (non-blocking) - UI sudah 'MandorApp'.
        msg = api_client.get(f"{base_url}/api/").json()["message"]
        if "TukangGaji" in msg:
            pytest.xfail(f"KNOWN: root API message masih '{msg}' (rebrand belum menyentuh backend root)")


# ---------- auth ----------

class TestAuth:
    def test_admin_login_success(self, api_client, base_url):
        resp = api_client.post(f"{base_url}/api/auth/admin-login", json={"username": "admin", "password": "tukangpro123"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["role"] == "admin"
        assert data["access_token"]

    def test_admin_login_wrong_password(self, api_client, base_url):
        resp = api_client.post(f"{base_url}/api/auth/admin-login", json={"username": "admin", "password": "salah123"})
        assert resp.status_code == 401

    def test_worker_login_success(self, api_client, base_url, login_worker):
        resp = api_client.post(f"{base_url}/api/auth/worker-login", json={"code": login_worker["code"], "pin": login_worker["pin"]})
        assert resp.status_code == 200
        data = resp.json()
        assert data["role"] == "worker"
        assert data["worker"]["code"] == login_worker["code"]
        assert data["worker"]["daily_rate"] == 160000
        assert data["worker"]["overtime_rate"] == 25000
        assert "pin_hash" not in data["worker"]

    def test_worker_login_wrong_pin(self, api_client, base_url, login_worker):
        wrong = "0000" if login_worker["pin"] != "0000" else "9999"
        resp = api_client.post(f"{base_url}/api/auth/worker-login", json={"code": login_worker["code"], "pin": wrong})
        assert resp.status_code == 401

    def test_worker_cannot_access_admin_endpoint(self, api_client, base_url, worker_token):
        resp = api_client.get(f"{base_url}/api/workers", headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 403

    def test_admin_cannot_access_worker_endpoint(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/worker/me")
        assert resp.status_code == 403

    def test_no_token_unauthorized(self, base_url):
        resp = requests.get(f"{base_url}/api/dashboard")
        assert resp.status_code == 401


# ---------- workers CRUD ----------

class TestWorkers:
    def test_create_worker_returns_code_and_pin(self, admin_client, base_url, test_worker):
        assert test_worker["code"].startswith("TG")
        assert test_worker["group"] == "Kuli"
        assert test_worker["daily_rate"] == 160000
        assert test_worker["overtime_rate"] == 25000
        get = admin_client.get(f"{base_url}/api/workers")
        persisted = [w for w in get.json() if w["id"] == test_worker["id"]][0]
        assert persisted["daily_rate"] == 160000
        assert persisted["overtime_rate"] == 25000

    def test_duplicate_name_rejected_409(self, admin_client, base_url, test_worker):
        resp = admin_client.post(f"{base_url}/api/workers", json={
            "name": "test pekerja pytest", "phone": "", "group": "Tukang",
            "daily_rate": 100000, "overtime_rate": 10000,
        })
        assert resp.status_code == 409
        assert "sudah terdaftar" in resp.json()["detail"].lower()

    def test_update_worker_rates(self, admin_client, base_url, test_worker):
        resp = admin_client.put(f"{base_url}/api/workers/{test_worker['id']}", json={"daily_rate": 170000, "overtime_rate": 22000})
        assert resp.status_code == 200
        data = resp.json()
        assert data["daily_rate"] == 170000
        assert data["overtime_rate"] == 22000
        get = admin_client.get(f"{base_url}/api/workers")
        persisted = [w for w in get.json() if w["id"] == test_worker["id"]][0]
        assert persisted["daily_rate"] == 170000
        # restore
        admin_client.put(f"{base_url}/api/workers/{test_worker['id']}", json={"daily_rate": 160000, "overtime_rate": 25000})

    def test_reset_pin_and_login(self, api_client, admin_client, base_url, test_worker):
        resp = admin_client.post(f"{base_url}/api/workers/{test_worker['id']}/reset-pin")
        assert resp.status_code == 200
        new_pin = resp.json()["pin"]
        assert len(new_pin) == 4
        login = api_client.post(f"{base_url}/api/auth/worker-login", json={"code": test_worker["code"], "pin": new_pin})
        assert login.status_code == 200

    def test_invalid_group_rejected(self, admin_client, base_url):
        resp = admin_client.post(f"{base_url}/api/workers", json={
            "name": "TEST Grup Ngawur", "phone": "", "group": "Sihir",
            "daily_rate": 100000, "overtime_rate": 10000,
        })
        assert resp.status_code == 422

    def test_missing_daily_rate_rejected(self, admin_client, base_url):
        resp = admin_client.post(f"{base_url}/api/workers", json={
            "name": "TEST Tanpa Tarif", "phone": "", "group": "Kuli", "overtime_rate": 10000,
        })
        assert resp.status_code == 422


# ---------- perhitungan gaji PER HARI + NO-REST SIANG 1 JAM ----------

class TestNoRestSiang1Jam:
    """Skenario pada TEST worker (harian 160.000, lembur 25.000/jam).

    ATURAN BARU: no-rest siang (11.30-12.30) = 1 jam berbayar x tarif lembur,
    BUKAN 1,5 jam. No-rest sore (17.00-18.00) tetap 1 jam.
    """

    def test_norest_siang_saja_rest_1_jam(self, admin_client, base_url, test_worker):
        # 07:00-17:00 + no_rest_siang -> rest_hours 1,0 & rest_pay 25.000 (BUKAN 1,5 j / 37.500)
        rec_id, calc = ensure_day(admin_client, base_url, test_worker["id"], "2026-09-20",
                                  "07:00", "17:00", no_siang=True)
        assert calc["day_type"] == "full"
        assert calc["base_pay"] == 160000
        assert calc["rest_hours"] == 1.0, f"no-rest siang harus 1,0 jam, dapat {calc['rest_hours']}"
        assert calc["rest_pay"] == 25000, f"rest_pay harus 25.000 (1 j x 25rb), dapat {calc['rest_pay']}"
        assert calc["total"] == 185000

    def test_norest_siang_sore_rest_2_jam(self, admin_client, base_url, test_worker):
        # 07:00-17:00 + no_rest_siang&sore -> rest_hours 2,0 & rest_pay 50.000
        rec_id, calc = ensure_day(admin_client, base_url, test_worker["id"], "2026-09-21",
                                  "07:00", "17:00", no_siang=True, no_sore=True)
        assert calc["day_type"] == "full"
        assert calc["rest_hours"] == 2.0, f"siang+sore harus 2,0 jam, dapat {calc['rest_hours']}"
        assert calc["rest_pay"] == 50000, f"rest_pay harus 50.000 (2 j x 25rb), dapat {calc['rest_pay']}"
        assert calc["total"] == 210000

    def test_norest_siang_butuh_shift_penuh(self, admin_client, base_url, test_worker):
        # 13:00-17:00 (setengah hari) + no_rest_siang -> TIDAK ada bonus siang (pagi tidak dikerjakan)
        rec_id, calc = ensure_day(admin_client, base_url, test_worker["id"], "2026-09-22",
                                  "13:00", "17:00", no_siang=True)
        assert calc["day_type"] == "half"
        assert calc["base_pay"] == 80000
        assert calc["rest_hours"] == 0
        assert calc["rest_pay"] == 0
        assert calc["total"] == 80000

    def test_full_day_norest_lembur_kombinasi(self, admin_client, base_url, test_worker):
        # 07:00-20:30 + no-rest siang&sore + lembur 2,5j -> 160rb + 50rb + 62,5rb = 272.500
        rec_id, calc = ensure_day(admin_client, base_url, test_worker["id"], "2026-09-24",
                                  "07:00", "20:30", no_siang=True, no_sore=True,
                                  lembur=True, lembur_hours=2.5, note="TEST kombinasi")
        assert calc["day_type"] == "full"
        assert calc["base_pay"] == 160000
        assert calc["rest_hours"] == 2.0
        assert calc["rest_pay"] == 50000
        assert calc["lembur_hours"] == 2.5
        assert calc["overtime_pay"] == 62500
        assert calc["total"] == 272500


class TestDailyWageRegression:
    """Regresi model gaji per hari (tidak berubah sesi ini)."""

    def test_half_day_morning(self, admin_client, base_url, test_worker):
        rec_id, calc = ensure_day(admin_client, base_url, test_worker["id"], "2026-09-23", "07:00", "11:45")
        assert calc["day_type"] == "half"
        assert calc["base_pay"] == 80000
        assert calc["total"] == 80000

    def test_full_day_plain(self, admin_client, base_url, test_worker):
        # hapus dulu bila test lain sudah memakai tanggal ini di modul yang sama
        rec_id, calc = ensure_day(admin_client, base_url, test_worker["id"], "2026-09-23", "07:00", "17:00")
        assert calc["day_type"] == "full"
        assert calc["base_pay"] == 160000
        assert calc["rest_hours"] == 0
        assert calc["rest_pay"] == 0
        assert calc["overtime_pay"] == 0
        assert calc["total"] == 160000


# ---------- dashboard ----------

class TestDashboard:
    def test_dashboard_today(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/dashboard")
        assert resp.status_code == 200
        data = resp.json()
        assert data["active_workers"] >= 3  # Bayu, Reza, Obeng (data asli)
        assert data["present"] >= 0
        groups = {g["group"] for g in data["groups"]}
        assert {"Besi", "Kayu"} <= groups


# ---------- payroll & GRAND TOTAL ----------

class TestPayroll:
    def test_payroll_grand_totals_match_rows(self, admin_client, base_url):
        """GET /api/payroll memuat total_days & total_overtime_hours yang konsisten
        dengan jumlah baris per pekerja (chip grand total di UI)."""
        resp = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2")
        assert resp.status_code == 200
        data = resp.json()
        assert data["start_date"] == "2026-09-16" and data["end_date"] == "2026-09-30"
        assert "total_days" in data and "total_overtime_hours" in data
        sum_days = sum(r["days"] + r["half_days"] for r in data["rows"])
        sum_ot = round(sum(r["overtime_hours"] for r in data["rows"]), 2)
        assert data["total_days"] == sum_days, f"total_days {data['total_days']} != sum rows {sum_days}"
        assert data["total_overtime_hours"] == sum_ot, f"total_overtime_hours {data['total_overtime_hours']} != sum rows {sum_ot}"
        assert data["total_gross"] == sum(r["gross"] for r in data["rows"])
        # data asli saat tes: Bayu 1h/3j, Reza 3h/9j, Obeng 1h/3j -> >= 5 hari, >= 15 j lembur
        assert data["total_days"] >= 5
        assert data["total_overtime_hours"] >= 15.0

    def test_payroll_row_fields_per_pekerja(self, admin_client, base_url):
        """Setiap baris memuat field untuk format UI 'N hari penuh · M 1/2 hari · Z j lembur'."""
        resp = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2")
        assert resp.status_code == 200
        for row in resp.json()["rows"]:
            for field in ("days", "half_days", "rest_hours", "overtime_hours", "base_pay", "rest_pay", "overtime_pay", "gross", "paid", "remaining"):
                assert field in row, f"Field '{field}' hilang di baris payroll"

    def test_payroll_group_filter(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2&group=Besi")
        assert resp.status_code == 200
        data = resp.json()
        assert data["rows"], "Filter grup Besi harus berisi Bayu & Reza"
        assert all(r["group"] == "Besi" for r in data["rows"])
        assert data["total_days"] == sum(r["days"] + r["half_days"] for r in data["rows"])

    def test_payment_history_endpoint(self, admin_client, base_url):
        resp = admin_client.get(f"{base_url}/api/payments?month={MONTH}&period=2")
        assert resp.status_code == 200
        assert isinstance(resp.json(), list)

    def test_mandor_payment_requires_name(self, admin_client, base_url, test_worker):
        resp = admin_client.post(f"{base_url}/api/payments", json={
            "month": MONTH, "period": 2, "method": "mandor", "mandor_name": "", "note": "",
            "items": [{"worker_id": test_worker["id"], "amount": 5000}],
        })
        assert resp.status_code == 422

    def test_create_payment_reduces_remaining_then_delete(self, admin_client, base_url, test_worker):
        wid = test_worker["id"]
        rows = admin_client.get(f"{base_url}/api/attendance?date=2026-09-23").json()["rows"]
        rec = [r for r in rows if r["worker"]["id"] == wid][0]["record"]
        if not rec:
            ensure_day(admin_client, base_url, wid, "2026-09-23", "07:00", "17:00")
        before = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        brow = [r for r in before["rows"] if r["code"] == test_worker["code"]][0]

        resp = admin_client.post(f"{base_url}/api/payments", json={
            "month": MONTH, "period": 2, "method": "langsung", "mandor_name": "", "note": "TEST bayar",
            "items": [{"worker_id": wid, "amount": 5000}],
        })
        assert resp.status_code == 201
        pid = resp.json()["id"]

        after = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        arow = [r for r in after["rows"] if r["code"] == test_worker["code"]][0]
        assert arow["paid"] == brow["paid"] + 5000
        assert arow["remaining"] == max(brow["gross"] - arow["paid"], 0)

        dele = admin_client.delete(f"{base_url}/api/payments/{pid}")
        assert dele.status_code == 200
        after2 = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        arow2 = [r for r in after2["rows"] if r["code"] == test_worker["code"]][0]
        assert arow2["paid"] == brow["paid"]


# ---------- worker self-service (TEST worker dengan PIN diketahui) ----------

class TestWorkerSelfService:
    def test_worker_payroll_period2_norest_1jam(self, api_client, base_url, worker_token):
        # TEST login worker: 23 Sep 07:00-17:00 + no-rest siang&sore -> 160rb + 50rb = 210rb
        resp = api_client.get(f"{base_url}/api/worker/payroll?month={MONTH}&period=2",
                              headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 200
        data = resp.json()
        assert data["days"] == 1 and data["half_days"] == 0
        assert data["rest_hours"] == 2.0, f"rest_hours harus 2,0 (siang 1j + sore 1j), dapat {data['rest_hours']}"
        assert data["rest_pay"] == 50000
        assert data["base_pay"] == 160000
        assert data["gross"] == 210000
        assert data["overtime_hours"] == 0

    def test_worker_attendance_month(self, api_client, base_url, worker_token):
        resp = api_client.get(f"{base_url}/api/worker/attendance?month={MONTH}",
                              headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 200
        rows = resp.json()["rows"]
        day_row = [r for r in rows if r["record"]["date"] == "2026-09-23"]
        assert day_row
        calc = day_row[0]["calc"]
        assert calc["day_type"] == "full"
        assert calc["rest_hours"] == 2.0
        assert calc["total"] == 210000

    def test_worker_me(self, api_client, base_url, worker_token, login_worker):
        resp = api_client.get(f"{base_url}/api/worker/me", headers={"Authorization": f"Bearer {worker_token}"})
        assert resp.status_code == 200
        assert resp.json()["code"] == login_worker["code"]
        assert "pin_hash" not in resp.json()


# ---------- Excel export + baris GRAND TOTAL ----------

class TestExport:
    def _load(self, resp):
        assert resp.status_code == 200
        assert resp.headers["Content-Type"].startswith("application/vnd.openxmlformats")
        assert resp.content[:2] == b"PK"
        return load_workbook(io.BytesIO(resp.content))

    def _get_export(self, base_url, kind, admin_token):
        """Unduh ekspor dengan retry (proxy kadang memotong stream -> BadZipFile)."""
        url = f"{base_url}/api/export/{kind}?month={MONTH}&period=2&token={admin_token}"
        for attempt in range(3):
            resp = requests.get(url, timeout=60)
            try:
                return self._load(resp)
            except Exception:
                if attempt == 2:
                    raise
        raise AssertionError("unreachable")

    def test_export_payroll_grand_total_row(self, base_url, admin_token, admin_client):
        """Baris terakhir ekspor gaji = GRAND TOTAL dengan jumlah hari, jam lembur,
        dan total rupiah yang cocok dengan jumlah baris pekerja & API payroll."""
        ws = self._get_export(base_url, "payroll", admin_token).active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        for col in ["Hari Penuh", "Setengah Hari", "Jam No-Rest", "Jam Lembur",
                    "Upah Harian (Rp)", "Upah No-Rest (Rp)", "Upah Lembur (Rp)", "Total (Rp)"]:
            assert col in headers, f"Kolom '{col}' tidak ada: {headers}"

        last = ws.max_row
        assert ws.cell(row=last, column=2).value == "GRAND TOTAL", \
            f"Baris terakhir (row {last}) bukan GRAND TOTAL: {ws.cell(row=last, column=2).value}"

        # jumlahkan baris pekerja (row 2 .. last-1) lalu bandingkan dengan GRAND TOTAL
        def col_sum(header):
            idx = headers.index(header) + 1
            return sum(ws.cell(row=r, column=idx).value or 0 for r in range(2, last))

        def gt(header):
            return ws.cell(row=last, column=headers.index(header) + 1).value

        assert gt("Hari Penuh") == col_sum("Hari Penuh")
        assert gt("Setengah Hari") == col_sum("Setengah Hari")
        assert gt("Jam No-Rest") == pytest.approx(col_sum("Jam No-Rest"))
        assert gt("Jam Lembur") == pytest.approx(col_sum("Jam Lembur"))
        assert gt("Total (Rp)") == col_sum("Total (Rp)")
        assert gt("Sisa (Rp)") == col_sum("Sisa (Rp)")

        # cross-check dengan API payroll (chip grand total di UI)
        api = admin_client.get(f"{base_url}/api/payroll?month={MONTH}&period=2").json()
        assert gt("Hari Penuh") + gt("Setengah Hari") == api["total_days"]
        assert gt("Jam Lembur") == pytest.approx(api["total_overtime_hours"])
        assert gt("Total (Rp)") == api["total_gross"]
        assert last - 2 == len(api["rows"])  # jumlah baris pekerja cocok

    def test_export_attendance_tipe_hari_column(self, base_url, admin_token):
        ws = self._get_export(base_url, "attendance", admin_token).active
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Tipe Hari" in headers, f"Kolom 'Tipe Hari' tidak ada: {headers}"
        tipe_idx = headers.index("Tipe Hari") + 1
        values = [ws.cell(row=r, column=tipe_idx).value for r in range(2, ws.max_row + 1)]
        assert "Penuh" in values, "Tidak ada baris 'Penuh' di ekspor absensi"

    def test_export_requires_admin_token(self, base_url, worker_token):
        resp = requests.get(f"{base_url}/api/export/payroll?month={MONTH}&period=2&token={worker_token}")
        assert resp.status_code == 401
